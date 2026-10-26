#!/usr/bin/env python3
"""
Build out/typemap.html: roll-up diagram of every Census building type -> Census category -> ownership -> dashboard segment.
Column letters are looked up in the raw workbooks; itemized coverage uses calendar 2025 NSA sums.
Usage: python3 typemap.py [raw_dir]
"""
import os, re, sys, html
import openpyxl
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "raw")
MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()

def load(f):
    ws = openpyxl.load_workbook(os.path.join(RAW, f), data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Date")
    h = [re.sub(r"\s+", " ", str(x).replace("_x000D_", "")).strip() if x else None for x in rows[hi]]
    d = {str(r[0])[:6]: r for r in rows[hi + 1:] if r and r[0] and re.match(r"[A-Z][a-z]{2}-\d\d", str(r[0]))}
    rel = next((re.search(r"Spending,\s+([A-Z][a-z]+ \d{1,2}, \d{4})", str(r[0])).group(1) for r in rows
                if r and r[0] and "Construction Spending," in str(r[0])), "")
    return {"h": h, "d": d, "rel": rel}
WB = {"P": load("privtime.xlsx"), "S": load("sltime.xlsx"), "U": load("pubtime.xlsx")}
TABNAME = {"P": "privtime", "S": "sltime", "U": "pubtime"}

def col(src, label, after=None):
    h = WB[src]["h"]; start = h.index(after) if after else 0
    return next(i for i in range(start, len(h)) if h[i] == label)
def year_sum(src, j, yy="25"):
    return sum(WB[src]["d"][f"{m}-{yy}"][j] for m in MONTHS)

# ---------------------------------------------------------------- hierarchy
# node: (label, census_label or None, children, flags)   flags: x = excluded
def N(label, census=None, kids=(), excluded=False, note=None):
    return {"label": label, "census": census or label, "kids": list(kids), "excluded": excluded, "note": note}

SEGMENTS = [
  ("Data centers", "data_center", [
     ("Private", "P", "Office", [N("Data center (in Office)", "Data center")])]),
  ("Health care", "health_care", [
     ("Private", "P", None, [N("Health care", kids=[N("Hospital"), N("Medical building"), N("Special care")])]),
     ("State & local", "S", None, [N("Health care", kids=[N("Hospital"), N("Medical building"), N("Special care")])]),
     ("Federal", "U", None, [N("Health care", "Federal Health care")])]),
  ("Aviation", "aviation", [
     ("Private", "P", "Transportation", [N("Air (all types)", "Air")]),
     ("State & local", "S", "Transportation", [N("Air", kids=[N("Air passenger terminal"), N("Runway", excluded=True)])]),
     ("Federal", "U", None, [N("Air (not published)", "Federal Transportation", excluded=True)])]),
  ("Educational", "educational", [
     ("Private", "P", None, [N("Educational", kids=[N("Preschool"), N("Primary/secondary"),
         N("Higher education", kids=[N("Instructional"), N("Dormitory"), N("Sports/recreation")]),
         N("Other educational", kids=[N("Gallery/museum")])])]),
     ("State & local", "S", None, [N("Educational", kids=[
         N("Primary/secondary", kids=[N("Elementary"), N("Middle/junior high"), N("High")]),
         N("Higher education", kids=[N("Instructional"), N("Dormitory"), N("Sports/recreation"), N("Infrastructure")]),
         N("Other educational", kids=[N("Library/archive")])])]),
     ("Federal", "U", None, [N("Educational", "Federal Educational")])]),
  ("Office (excl. data centers)", "office_ex_dc", [
     ("Private", "P", None, [N("Office", kids=[N("General"), N("Financial")])]),
     ("State & local", "S", None, [N("Office")]),
     ("Federal", "U", None, [N("Office", "Federal Office")])]),
  ("Commercial (incl. farm)", "commercial", [
     ("Private", "P", None, [N("Commercial (incl. farm)", "Commercial (inc. Farm)", kids=[
         N("Automotive", kids=[N("Sales"), N("Service/parts"), N("Parking")]),
         N("Food/beverage", kids=[N("Food"), N("Dining/drinking")]),
         N("Multi-retail", kids=[N("General merchandise"), N("Shopping center"), N("Shopping mall")]),
         N("Other commercial", kids=[N("Drug store"), N("Building supply store"), N("Other stores")]),
         N("Warehouse", kids=[N("General commercial"), N("Mini-storage")])])]),
     ("State & local", "S", None, [N("Commercial", kids=[N("Automotive", kids=[N("Parking")])])]),
     ("Federal", "U", None, [N("Commercial", "Federal Commercial")])]),
  ("Manufacturing", "manufacturing", [
     ("Private", "P", None, [N("Manufacturing", kids=[N("Food/beverage/tobacco", "Food/beverage/ tobacco"), N("Chemical"),
         N("Plastic/rubber"), N("Nonmetallic mineral"), N("Fabricated metal"),
         N("Computer/electronic/electrical", "Computer/ electronic/ electrical"), N("Transportation equipment")])])]),
  ("Amusement & recreation", "amusement", [
     ("Private", "P", None, [N("Amusement & recreation", "Amusement and Recreation", kids=[
         N("Theme/amusement park"), N("Sports"), N("Fitness"), N("Performance/meeting center", "Performance/ meeting center"),
         N("Social center"), N("Movie theater/studio", "Movie theater/ studio")])]),
     ("State & local", "S", None, [N("Amusement & recreation", "Amusement and recreation", kids=[
         N("Sports"), N("Performance/meeting center", "Performance/ meeting center", kids=[N("Convention center")]),
         N("Social center", kids=[N("Neighborhood center")]), N("Park/camp")])]),
     ("Federal", "U", None, [N("Amusement & recreation", "Federal Amusement and recreation")])]),
  ("Lodging", "lodging", [("Private", "P", None, [N("Lodging")])]),
  ("Public safety", "public_safety", [
     ("State & local", "S", None, [N("Public safety", "Public Safety", kids=[
         N("Correctional", kids=[N("Detention"), N("Police/sheriff")]),
         N("Other public safety", kids=[N("Fire/rescue")])])]),
     ("Federal", "U", None, [N("Public safety", "Federal Public safety")])]),
  ("Religious", "religious", [
     ("Private", "P", None, [N("Religious", kids=[N("House of worship"), N("Other religious", kids=[N("Auxiliary building")])])])]),
]
FOCUS = {"data_center": "--s1", "health_care": "--s2", "aviation": "--s3", "educational": "--s7"}

# resolve column letters + coverage
coverage = []
def resolve(node, src, after):
    j = col(src, node["census"], after)
    node["ref"] = L(j + 1); node["j"] = j
    for k in node["kids"]:
        resolve(k, src, node["census"])
    if node["kids"] and src in ("P", "S"):
        pv = year_sum(src, j); kv = sum(year_sum(src, k["j"]) for k in node["kids"])
        node["cov"] = kv / pv if pv else None
        if node["cov"] is not None and node["cov"] < 0.95:
            coverage.append((TABNAME[src], node["label"], node["cov"]))
for seg, sid, owns in SEGMENTS:
    for own, src, after, nodes in owns:
        for n in nodes:
            resolve(n, src, after)

# ---------------------------------------------------------------- layout
COLX = {0: 0, 1: 214, 2: 428, 3: 618}   # sub-type | category | ownership | segment  (left -> right)
COLW = {0: 196, 1: 196, 2: 172, 3: 196}
PITCH, NH, SEGGAP, TOP = 24, 18, 22, 34
y = TOP
placed = []
def place2(node, level):
    global y
    if not node["kids"]:
        node["cy"] = y; y += PITCH
    elif level >= 1:
        # group label in the sub-type column: its own row, children listed beneath it
        node["cy"] = y; y += PITCH
        for k in node["kids"]: place2(k, level + 1)
    else:
        ys = [place2(k, level + 1) for k in node["kids"]]
        node["cy"] = (ys[0] + ys[-1]) / 2
    node["level"] = level
    placed.append(node)
    return node["cy"]

# Column assignment: level 0 (category) -> col 1; level >= 1 -> col 0. A level-2 node (e.g. Higher education > Instructional)
# is drawn in col 0, indented under its level-1 parent, which sits on its own row as a group label.
SEGS = []
for seg, sid, owns in SEGMENTS:
    seg_start = y
    owners = []
    for own, src, after, nodes in owns:
        tops = [place2(n, 0) for n in nodes]
        owners.append({"own": own, "src": src, "after": after, "nodes": nodes, "cy": (tops[0] + tops[-1]) / 2})
    SEGS.append({"seg": seg, "sid": sid, "owners": owners, "cy": (owners[0]["cy"] + owners[-1]["cy"]) / 2,
                 "y0": seg_start, "y1": y - PITCH})
    y += SEGGAP
H = y + 10
W = COLX[3] + COLW[3] + 4

# ---------------------------------------------------------------- SVG
def esc(s): return html.escape(s, quote=True)
out = []
def rect_node(x, cy, w, label, ref=None, cls="", title=None, indent=0):
    yy = cy - NH / 2
    t = f"<title>{esc(title)}</title>" if title else ""
    out.append(f'<g class="nd {cls}">{t}<rect x="{x+indent:.1f}" y="{yy:.1f}" width="{w-indent:.1f}" height="{NH}" rx="4"/>'
               f'<text x="{x+indent+8:.1f}" y="{cy+4:.1f}" class="lb">{esc(label)}</text>'
               + (f'<text x="{x+w-7:.1f}" y="{cy+3.5:.1f}" class="rf" text-anchor="end">{esc(ref)}</text>' if ref else "")
               + "</g>")
def elbow(x1, y1, x2, y2, cls=""):
    xm = x2 - 12
    out.append(f'<path class="eg {cls}" d="M{x1:.1f},{y1:.1f} H{xm:.1f} V{y2:.1f} H{x2-1:.1f}"/>')

# column headers
for c, lab in [(0, "Census sub-type"), (1, "Census category"), (2, "Ownership"), (3, "Dashboard segment")]:
    out.append(f'<text x="{COLX[c]}" y="14" class="ch">{lab.upper()}</text>')

for S in SEGS:
    sid = S["sid"]; tone = "focus" if sid in FOCUS else "plain"
    style = f' style="--seg:var({FOCUS[sid]})"' if sid in FOCUS else ""
    out.append(f'<g class="segblock {tone}"{style}>')
    # band
    out.append(f'<rect class="band" x="-6" y="{S["y0"]-PITCH/2-3:.1f}" width="{W+12}" height="{S["y1"]-S["y0"]+PITCH+6:.1f}" rx="6"/>')
    for O in S["owners"]:
        src = O["src"]; own = O["own"]
        ocls = "own priv" if own == "Private" else "own pub"
        for n in O["nodes"]:
            def draw(node, parent_x, parent_cy, level):
                ex = node["excluded"]
                cls = ("cat" if level == 0 else "sub") + (" ex" if ex else "")
                if level == 0:
                    x, w, indent = COLX[1], COLW[1], 0
                else:
                    x, w, indent = COLX[0], COLW[0], 12 * (level - 1)
                tip = f'{TABNAME[src]} column {node["ref"]}: "{node["census"]}"'
                lab = node["label"] + (" (excluded)" if ex and "not published" not in node["label"] else "")
                # edge to parent
                if level == 0:
                    elbow(x + w, node["cy"], parent_x, parent_cy, "ex" if ex else "")
                else:
                    # sub-types: bracket from left of parent? draw vertical spine at parent's left side within col0/col1
                    if level == 1:
                        elbow(x + indent + (w - indent), node["cy"], parent_x, parent_cy, "ex" if ex else "")
                    else:
                        xs = x + 12 * (level - 2) + 5
                        out.append(f'<path class="eg sp{" ex" if ex else ""}" d="M{x+indent:.1f},{node["cy"]:.1f} H{xs:.1f} V{parent_cy+NH/2:.1f}"/>')
                ref = None if "not published" in node["label"] else node["ref"]
                if ref is None: tip = "Federal air spending is not published separately (only federal transportation, pubtime column " + node["ref"] + ")"
                rect_node(x, node["cy"], w, lab, ref, cls, tip, indent)
                for k in node["kids"]:
                    if level == 0:
                        draw(k, x, node["cy"], 1)
                    else:
                        draw(k, x + indent, node["cy"], level + 1)
            draw(n, COLX[2], O["cy"], 0)
        oref = {"P": "privtime", "S": "sltime", "U": "pubtime"}[src]
        olab = {"Private": "Private", "State & local": "Public · state & local", "Federal": "Public · federal"}[own]
        rect_node(COLX[2], O["cy"], COLW[2], olab, None, ocls, f"Source tab: {oref}")
        elbow(COLX[2] + COLW[2], O["cy"], COLX[3], S["cy"])
    rect_node(COLX[3], S["cy"], COLW[3], S["seg"], None, "segn", None)
    out.append("</g>")

defs = ('<defs><pattern id="hatch" patternUnits="userSpaceOnUse" width="6" height="6" patternTransform="rotate(45)">'
        '<line x1="0" y1="0" x2="0" y2="6" class="hl"/></pattern>'
        '<marker id="arr" viewBox="0 0 8 8" refX="7" refY="4" markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
        '<path d="M0,0 L8,4 L0,8 z" class="ah"/></marker></defs>')
n_types = sum(1 for p in placed if not p["kids"])
svg = (f'<svg viewBox="-8 0 {W+16} {H}" width="{W+16}" role="img" aria-label="Roll-up of {n_types} Census building types into '
       f'ownership and the eleven dashboard segments">{defs}{"".join(out)}</svg>')

# ---------------------------------------------------------------- excluded list + coverage notes
EXCL = [
  ("Private", ["Transportation › Land", "Communication", "Power (incl. gas and oil) › Electric",
               "Public safety, highway and street, sewage and waste, water supply, conservation (in the private total, not itemized)"]),
  ("Public · state & local", ["Transportation › Air › Runway", "Transportation › Land (land passenger terminal, mass transit)",
               "Transportation › Water (dock/marina)", "Power", "Highway and street (pavement, lighting, bridge, rest facility)",
               "Sewage and waste disposal", "Water supply", "Conservation and development"]),
  ("Public · federal", ["Transportation (incl. air)", "Power", "Highway and street", "Conservation and development"]),
  ("All owners", ["Residential", "Categories Census does not itemize for an owner: public lodging, religious, communication and manufacturing"]),
]
excl_html = "".join(f'<div class="xg"><h3>{esc(g)}</h3><ul>' + "".join(f"<li>{esc(i)}</li>" for i in items) + "</ul></div>" for g, items in EXCL)
cov_html = "".join(f"<li><b>{esc(lab)}</b> <span class=\"mono\">{tab}</span>: {c:.0%}</li>"
                   for tab, lab, c in sorted(coverage, key=lambda x: x[2]))

page = f"""<title>C30 Building Type Map</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Public+Sans:wght@400;500;600;700;800&family=IBM+Plex+Mono:wght@400;500&display=swap">
<style>
/* Layout: one wide roll-up diagram (types -> category -> ownership -> segment), then exclusions and coverage notes.
   Same tokens as the C30 dashboard: blueprint-blue ink on cool grey, segment hues for the four highlighted segments. */
:root{{
  --bg:#f3f4f2; --surface:#fcfcfb; --surface-2:#eef0ec; --ink:#121417; --ink-2:#4b4f55; --muted:#7e8288;
  --line:#c9ccc6; --edge:#9aa0a6; --accent:#1f4e8c;
  --s1:#2a78d6; --s2:#eb6834; --s3:#1baf7a; --s7:#4a3aa7;
  --f-ui:"Public Sans",system-ui,-apple-system,"Segoe UI",sans-serif;
  --f-mono:"IBM Plex Mono",ui-monospace,Menlo,Consolas,monospace;
}}
@media (prefers-color-scheme: dark){{:root:not([data-theme="light"]){{
  --bg:#0f1113; --surface:#17191c; --surface-2:#1e2124; --ink:#f1f2ef; --ink-2:#c3c5bf; --muted:#8b8f93;
  --line:#3b3f43; --edge:#6b7076; --accent:#7fa8e6; --s1:#3987e5; --s2:#d95926; --s3:#199e70; --s7:#9085e9; color-scheme:dark}}}}
:root[data-theme="dark"]{{
  --bg:#0f1113; --surface:#17191c; --surface-2:#1e2124; --ink:#f1f2ef; --ink-2:#c3c5bf; --muted:#8b8f93;
  --line:#3b3f43; --edge:#6b7076; --accent:#7fa8e6; --s1:#3987e5; --s2:#d95926; --s3:#199e70; --s7:#9085e9; color-scheme:dark}}
*{{box-sizing:border-box}}
body{{background:var(--bg);color:var(--ink);font-family:var(--f-ui);font-size:14px;line-height:1.45;margin:0}}
.wrap{{max-width:1120px;margin:0 auto;padding-inline:20px;padding-block:24px 56px;display:grid;gap:18px}}
h1{{font-size:24px;font-weight:800;margin:0;text-wrap:balance}} h2{{font-size:16px;margin:0}} h3{{font-size:12px;margin:0 0 4px;text-transform:uppercase;letter-spacing:.06em;color:var(--muted)}}
p{{margin:0;max-width:78ch}} .ink2{{color:var(--ink-2)}} .mono{{font-family:var(--f-mono);font-size:12px}}
.eyebrow{{font-size:11px;font-weight:600;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}}
.panel{{background:var(--surface);border:1px solid var(--line);border-radius:8px;padding:16px;display:grid;gap:12px;min-width:0}}
.legend{{display:flex;flex-wrap:wrap;gap:8px 18px;font-size:12.5px;color:var(--ink-2);align-items:center}}
.legend span{{display:inline-flex;align-items:center;gap:7px}}
.sw{{width:22px;height:13px;border-radius:3px;border:1px solid var(--edge);display:inline-block;background:var(--surface)}}
.sw.pub{{background:repeating-linear-gradient(45deg,var(--surface) 0 3px,var(--line) 3px 4.5px)}}
.sw.ex{{border-style:dashed;opacity:.7}}
.dot{{width:11px;height:11px;border-radius:3px;display:inline-block}}
figure{{margin:0;display:grid;gap:8px;min-width:0}}
.scroll{{overflow-x:auto;-webkit-overflow-scrolling:touch}}
.scroll svg{{display:block;max-width:none;height:auto;min-width:860px;width:100%}}
figcaption{{font-size:12.5px;color:var(--ink-2)}}
/* svg classes */
.ch{{font:600 10.5px var(--f-ui);letter-spacing:.08em;fill:var(--muted)}}
.band{{fill:transparent;stroke:none}}
.segblock.focus .band{{fill:color-mix(in oklab,var(--seg) 6%,transparent)}}
.nd rect{{fill:var(--surface);stroke:var(--line);stroke-width:1}}
.nd .lb{{font:500 11.5px var(--f-ui);fill:var(--ink)}}
.nd .rf{{font:400 9.5px var(--f-mono);fill:var(--muted)}}
.nd.cat rect{{stroke:var(--edge)}}
.nd.cat .lb{{font-weight:600}}
.nd.sub .lb{{fill:var(--ink-2)}}
.nd.own rect{{stroke:var(--edge)}}
.nd.own.pub rect{{fill:url(#hatch)}}
.nd.own .lb{{font-weight:600}}
.hl{{stroke:var(--line);stroke-width:2}}
.nd.segn rect{{fill:var(--ink);stroke:none}}
.nd.segn .lb{{fill:var(--surface);font-weight:700;font-size:12.5px}}
.segblock.focus .nd.segn rect{{fill:var(--seg)}}
.segblock.focus .nd.segn .lb{{fill:#fff}}
.nd.ex rect{{stroke-dasharray:3 3;fill:transparent}}
.nd.ex .lb,.nd.ex .rf{{fill:var(--muted)}}
.eg{{fill:none;stroke:var(--edge);stroke-width:1.2}}
.segblock.focus .eg{{stroke:color-mix(in oklab,var(--seg) 70%,var(--edge))}}
.eg.ex{{stroke-dasharray:3 3;stroke:var(--muted)!important;opacity:.7}}
.ah{{fill:var(--edge)}}
.grid2{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:14px}}
@media (max-width:900px){{.grid2{{grid-template-columns:repeat(2,minmax(0,1fr))}}}}
@media (max-width:520px){{.grid2{{grid-template-columns:1fr}}}}
.xg ul,.cov{{margin:0;padding-left:18px;font-size:12.5px;color:var(--ink-2);display:grid;gap:3px}}
.cov{{columns:2;display:block}} .cov li{{margin-bottom:3px;break-inside:avoid}}
@media (max-width:640px){{.cov{{columns:1}}}}
</style>
<div class="wrap">
  <header style="display:grid;gap:8px">
    <div class="eyebrow">US Census Bureau · Value of Construction Put in Place (C30) · release {esc(WB["P"]["rel"])}</div>
    <h1>C30 Building Type Map</h1>
    <p class="ink2">Every Census building type in the dashboard's scope, read left to right: sub-type → Census category → ownership → the dashboard segment it rolls into. {n_types} types feed eleven segments; the four highlighted segments are shaded.</p>
    <div class="legend">
      <span><i class="sw"></i>Private</span>
      <span><i class="sw pub"></i>Public (state &amp; local or federal)</span>
      <span><i class="sw ex"></i>Excluded from buildings</span>
      <span><i class="dot" style="background:var(--s1)"></i>Data centers</span>
      <span><i class="dot" style="background:var(--s2)"></i>Health care</span>
      <span><i class="dot" style="background:var(--s3)"></i>Aviation</span>
      <span><i class="dot" style="background:var(--s7)"></i>Educational</span>
    </div>
  </header>
  <section class="panel">
    <figure>
      <div class="scroll">{svg}</div>
      <figcaption>Letters on each box are the Census column in that owner's workbook: <span class="mono">privtime</span> for private, <span class="mono">sltime</span> for state &amp; local, <span class="mono">pubtime</span> for federal. Hover a box for the full source reference. Federal spending is published only at category level. In the dashboard, public totals for office, commercial and amusement come from <span class="mono">pubtime</span>, which equals state &amp; local plus federal.</figcaption>
    </figure>
  </section>
  <section class="panel">
    <h2>Excluded from buildings</h2>
    <p class="ink2" style="font-size:13px">Census categories left out because they are infrastructure or not buildings.</p>
    <div class="grid2">{excl_html}</div>
  </section>
  <section class="panel">
    <h2>Categories not fully itemized</h2>
    <p class="ink2" style="font-size:13px">Share of each category's 2025 spending that Census breaks out into the sub-types shown. The rest is published only in the category total; the dashboard reports it as "Other / not shown separately".</p>
    <ul class="cov">{cov_html}</ul>
  </section>
</div>
"""
os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
open(os.path.join(HERE, "out", "typemap.html"), "w").write(page)
print("types", n_types, "segments", len(SEGS), "height", H, "coverage notes", len(coverage))
