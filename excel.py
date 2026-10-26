#!/usr/bin/env python3
"""
Build C30_Building_Dashboard.xlsx: a formula-driven Excel version of the dashboard.

Every number on the Dashboard is a formula chain back to the six raw Census tabs:
  raw tabs (pasted Census sheets) -> Map (column lookup by header label) -> Pull_NSA / Pull_SA
  (one column per Census series, row found by month name) -> Segments (signed sums from segments.json)
  -> Dashboard / ChartData.
Lookups match month labels ("Aug-26*") and header labels, so pasting a newer Census sheet over a raw
tab (rows shift down each month) updates everything.

Usage: python3 excel.py [raw_dir] [out.xlsx]
"""
import os, sys, json, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "raw")
OUTF = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "out", "C30_Building_Dashboard.xlsx")
CFG = json.load(open(os.path.join(HERE, "segments.json")))

START = datetime.date(2014, 1, 1)          # first month pulled (data-center series begins Jan 2014)
NROWS = 264                                # 2014-01 .. 2035-12
R0 = 9                                     # first data row on Pull_* / Segments
RN = R0 + NROWS - 1
MONTHS_ARR = '{"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"}'
CHOOSE_M = 'CHOOSE(MONTH({c}),"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec")'
RAWRNG = "$A$1:$DZ$900"

# ---------------------------------------------------------------- styling
ARIAL = "Arial"
INK, MUTED, GRID = "1F2328", "6B7078", "D9DCD6"
NAVY = "1F4E8C"
FOCUS = {"data_center": "2A78D6", "health_care": "EB6834", "aviation": "1BAF7A", "educational": "4A3AA7"}
GRAYS = ["4B4F55", "A2A5A9", "6E7277", "C3C5C8", "35383C", "8A8E93", "D6D8DA"]
def F(**k): k.setdefault("name", ARIAL); k.setdefault("size", 10); return Font(**k)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=GRID)
BOT = Border(bottom=thin)
HDR_FILL = fill("EEF0EC")
GREEN_LINK = "008000"
FMT_M = '#,##0;(#,##0);"-"'
FMT_B = '$#,##0.0,"B";($#,##0.0,"B");"-"'
FMT_DB = '+$#,##0.0,"B";-$#,##0.0,"B";"-"'
FMT_P = '+0.0%;-0.0%;0.0%'
FMT_SH = '0.0%'

def shade(hexc, t):
    """Blend hex colour toward white by t (0 = colour, 1 = white)."""
    r, g, b = (int(hexc[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{int(v + (255 - v) * t):02X}" for v in (r, g, b))

wb = openpyxl.Workbook()
wb._named_styles["Normal"].font = Font(name=ARIAL, size=10)
first = wb.active

def sheet(name, color=None):
    ws = wb.create_sheet(name)
    if color: ws.sheet_properties.tabColor = color
    ws.sheet_view.showGridLines = False
    return ws

def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)

# ---------------------------------------------------------------- tabs & components
TABS = []   # (tabname, sourcekey, adj)
for key, src in CFG["sources"].items():
    for adj in ("nsa", "sa"):
        if src.get(adj): TABS.append((src[adj].replace(".xlsx", ""), key, adj))
TABORDER = ["privtime", "pubtime", "sltime", "privsatime", "pubsatime", "slsatime"]
TABS.sort(key=lambda t: TABORDER.index(t[0]) if t[0] in TABORDER else 99)
TAB_OF = {(k, a): t for t, k, a in TABS}
SECTOR_LABEL = {"private": "Private", "public": "Public", "statelocal": "State & local"}

def pattern(source, col):
    """Census header label -> MATCH pattern. Public workbook headers carry a line break after 'Public'."""
    for pre in ("Public ", "State ", "Federal "):
        if source == "public" and col.startswith(pre):
            return pre.strip() + "*" + col[len(pre):]
    return col

src = lambda c: c.get("source", c["sector"])
comp_keys = []
def addc(s, col):
    if (s, col) not in comp_keys: comp_keys.append((s, col))
for s in CFG["segments"]:
    for c in s["components"]: addc(src(c), c["col"])
    for ss in s["subsegments"] + s.get("breakdown", []):
        for c in ss["components"]: addc(src(c), c["col"])
CHECK_COMPS = [("private", "Nonresidential"), ("public", "Public Nonresidential"),
               ("public", "Federal Health care"), ("public", "Federal Educational")]
for k in CHECK_COMPS: addc(*k)
SEG_COMPS = [k for k in comp_keys if k not in CHECK_COMPS or any(
    (src(c), c["col"]) == k for s in CFG["segments"] for c in s["components"])]

# ---------------------------------------------------------------- raw tabs (exact copies of Census sheets)
raw_meta = {}
for tab, key, adj in TABS:
    path = os.path.join(RAW, tab + ".xlsx")
    srcws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    ws = wb.create_sheet(tab); ws.sheet_properties.tabColor = "9AA0A6"
    for row in srcws.iter_rows():
        for c in row:
            if c.value is not None:
                v = c.value.replace("_x000D_", "") if isinstance(c.value, str) else c.value
                ws.cell(c.row, c.column, v)
    ws.column_dimensions["A"].width = 11
    raw_meta[tab] = srcws.title

# ---------------------------------------------------------------- Checks (control cells + reconciliation)
CK = sheet("Checks", "C9A227")
CK["A1"] = "Checks and control cells"; CK["A1"].font = F(size=14, bold=True)
CK["A2"] = ("Every other tab reads the latest month and release from this sheet. All six raw tabs must show the same "
            "latest month and release; otherwise the status cell turns red.")
CK["A2"].font = F(color=MUTED)
hdr = ["Raw tab", "Census sheet name", "Header row (\"Date\")", "Latest month label", "Latest month",
       "Release (from footer)", "Same month as privtime", "Same release as privtime"]
for j, h in enumerate(hdr, 1):
    c = CK.cell(4, j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
    c.alignment = Alignment(wrap_text=True, vertical="bottom")
TABROW = {}
for i, (tab, key, adj) in enumerate(TABS):
    r = 5 + i; TABROW[tab] = r
    CK.cell(r, 1, tab).font = F(bold=True)
    CK.cell(r, 2, raw_meta[tab]).font = F(color=MUTED)
    CK.cell(r, 3, f'=MATCH("Date",{tab}!$A:$A,0)')
    CK.cell(r, 4, f'=INDEX({tab}!$A:$A,C{r}+1)')
    CK.cell(r, 5, f'=DATE(2000+VALUE(MID(D{r},5,2)),MATCH(LEFT(D{r},3),{MONTHS_ARR},0),1)').number_format = "mmm yyyy"
    CK.cell(r, 6, f'=IFERROR(MID(INDEX({tab}!$A:$A,MATCH("Source: U.S. Census Bureau*",{tab}!$A:$A,0)),'
                  f'FIND("Spending, ",INDEX({tab}!$A:$A,MATCH("Source: U.S. Census Bureau*",{tab}!$A:$A,0)))+10,'
                  f'FIND(". ",INDEX({tab}!$A:$A,MATCH("Source: U.S. Census Bureau*",{tab}!$A:$A,0)),'
                  f'FIND("Spending, ",INDEX({tab}!$A:$A,MATCH("Source: U.S. Census Bureau*",{tab}!$A:$A,0))))'
                  f'-FIND("Spending, ",INDEX({tab}!$A:$A,MATCH("Source: U.S. Census Bureau*",{tab}!$A:$A,0)))-10),"not found")')
    CK.cell(r, 7, f'=IF(E{r}=$E$5,"OK","MISMATCH")')
    CK.cell(r, 8, f'=IF(F{r}=$F$5,"OK","MISMATCH")')
rS = 5 + len(TABS) + 1
CK.cell(rS, 1, "Status").font = F(bold=True)
CK.cell(rS, 2, f'=IF(COUNTIF(G5:H{rS-2},"MISMATCH")=0,"All raw tabs consistent","Raw tabs out of sync: re-paste the tabs marked MISMATCH")')
CK.cell(rS, 2).font = F(bold=True)
CK.conditional_formatting.add(f"G5:H{rS-2}", CellIsRule(operator="equal", formula=['"MISMATCH"'], font=F(color="B42F2F", bold=True), fill=fill("FBE3E3")))
CK.conditional_formatting.add(f"B{rS}", FormulaRule(formula=[f'COUNTIF($G$5:$H${rS-2},"MISMATCH")>0'], font=F(color="B42F2F", bold=True), fill=fill("FBE3E3")))
CK.conditional_formatting.add(f"B{rS}", FormulaRule(formula=[f'COUNTIF($G$5:$H${rS-2},"MISMATCH")=0'], font=F(color="006300", bold=True), fill=fill("E3F4E3")))

rc = rS + 2
CK.cell(rc, 1, "Control cells (named ranges used throughout)").font = F(size=12, bold=True)
ctrl = [("LatestLabel", "Latest month label", "=D5", None),
        ("LatestMonth", "Latest month", "=E5", "mmm yyyy"),
        ("LY", "Current year", "=YEAR(LatestMonth)", "0"),
        ("LM", "Months in YTD", "=MONTH(LatestMonth)", "0"),
        ("ReleaseText", "Census release", "=F5", None),
        ("LatestFlag", "Latest month status", '=IF(RIGHT(LatestLabel,1)="p","preliminary",IF(RIGHT(LatestLabel,1)="r","revised","final"))', None),
        ("FirstFullYear", "First year of 10-year history", "=LY-10", "0"),
        ("LastFullYear", "Last full year", "=LY-1", "0")]
for i, (nm, lab, fml, fmt) in enumerate(ctrl):
    r = rc + 1 + i
    CK.cell(r, 1, nm).font = F(color=MUTED)
    CK.cell(r, 2, lab)
    c = CK.cell(r, 3, fml); c.font = F(bold=True)
    if fmt: c.number_format = fmt
    name(nm, f"Checks!$C${r}")
CK_RECON = rc + len(ctrl) + 3
CK.column_dimensions["A"].width = 26; CK.column_dimensions["B"].width = 42
for col in "CDEFGH": CK.column_dimensions[col].width = 18

# ---------------------------------------------------------------- Map (Census column located by header label)
MP = sheet("Map", "5B8DD6")
MP["A1"] = "Map: where each Census series sits in the raw tabs"; MP["A1"].font = F(size=14, bold=True)
MP["A2"] = ("Columns are found by matching the header label in each raw tab's \"Date\" row, so the workbook keeps "
            "working if Census inserts or moves columns. An error in column G means Census renamed a header: "
            "update the pattern in column E.")
MP["A2"].font = F(color=MUTED)
mh = ["ID", "Raw tab", "Sector", "Adjustment", "Match pattern (header label)", "Census label", "Column #", "Column", "Used for"]
for j, h in enumerate(mh, 1):
    c = MP.cell(4, j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
usage = {}
for s in CFG["segments"]:
    for c in s["components"]: usage.setdefault((src(c), c["col"]), set()).add(s["name"])
    for ss in s["subsegments"] + s.get("breakdown", []):
        for c in ss["components"]: usage.setdefault((src(c), c["col"]), set()).add(s["name"] + " detail")
for k in CHECK_COMPS: usage.setdefault(k, set()).add("Checks")
MAPROW = {}   # (source, col, adj) -> row
r = 5
for adj in ("nsa", "sa"):
    for (s_, col) in comp_keys:
        tab = TAB_OF[(s_, adj)]
        cid = f"{tab}:{col}"
        MAPROW[(s_, col, adj)] = r
        MP.cell(r, 1, cid).font = F(color=MUTED, size=9)
        MP.cell(r, 2, tab)
        MP.cell(r, 3, SECTOR_LABEL[s_])
        MP.cell(r, 4, "Not seasonally adjusted" if adj == "nsa" else "SAAR")
        c = MP.cell(r, 5, pattern(s_, col)); c.font = F(color="0000FF")
        MP.cell(r, 6, col)
        MP.cell(r, 7, f'=MATCH(E{r},INDEX({tab}!{RAWRNG},Checks!$C${TABROW[tab]},0),0)')
        MP.cell(r, 8, f'=SUBSTITUTE(ADDRESS(1,G{r},4),"1","")')
        MP.cell(r, 9, ", ".join(sorted(usage.get((s_, col), [])))).font = F(color=MUTED)
        r += 1
for col, w in zip("ABCDEFGHI", [30, 11, 12, 22, 34, 34, 10, 8, 50]):
    MP.column_dimensions[col].width = w
MP.freeze_panes = "A5"

# ---------------------------------------------------------------- Pull_NSA / Pull_SA
PULLCOL = {}   # (source, col, adj) -> column letter on Pull sheet
def build_pull(adj):
    ws = sheet("Pull_" + adj.upper(), "5B8DD6")
    tabs = [t for t, k, a in TABS if a == adj]
    ws["A1"] = f"Pull_{adj.upper()}: every Census series used, one row per month ({'not seasonally adjusted' if adj=='nsa' else 'seasonally adjusted annual rate'}, $ millions)"
    ws["A1"].font = F(size=13, bold=True)
    ws["A2"] = ("Each value = INDEX(raw tab, row of that month, column from Map). Row numbers in columns D onward and the "
                "Column row above each series give the exact source cell (e.g. column J + row 5 = cell J5).")
    ws["A2"].font = F(color=MUTED)
    labels = ["Series ID", "Raw tab", "Census label", "Column #", "Column"]
    for i, lab in enumerate(labels):
        ws.cell(3 + i, 3, lab).font = F(bold=True, color=MUTED)
        ws.cell(3 + i, 3).alignment = Alignment(horizontal="right")
    heads = ["Month", "Match pattern", "Year"] + [f"Row in {t}" for t in tabs]
    for j, h in enumerate(heads, 1):
        c = ws.cell(8, j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
        c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center")
    rowcol = {t: L(4 + i) for i, t in enumerate(tabs)}
    c0 = 4 + len(tabs)
    for i in range(NROWS):
        r = R0 + i
        ws.cell(r, 1, f"=DATE({START.year},{START.month},1)" if i == 0 else f"=EDATE(A{r-1},1)").number_format = "mmm yyyy"
        ws.cell(r, 2, f'={CHOOSE_M.format(c=f"A{r}")}&"-"&TEXT(MOD(YEAR(A{r}),100),"00")&"*"').font = F(color=MUTED, size=9)
        ws.cell(r, 3, f"=YEAR(A{r})").font = F(color=MUTED)
        for t in tabs:
            ws[f"{rowcol[t]}{r}"] = f'=IF(A{r}>LatestMonth,"",MATCH(B{r},{t}!$A:$A,0))'
            ws[f"{rowcol[t]}{r}"].font = F(color=MUTED, size=9)
    for j, (s_, col) in enumerate(comp_keys):
        cl = L(c0 + j); tab = TAB_OF[(s_, adj)]; mr = MAPROW[(s_, col, adj)]
        PULLCOL[(s_, col, adj)] = cl
        ws[f"{cl}3"] = f"{tab}:{col}"; ws[f"{cl}4"] = tab; ws[f"{cl}5"] = col
        ws[f"{cl}6"] = f"=Map!G{mr}"; ws[f"{cl}7"] = f"=Map!H{mr}"
        for rr in range(3, 8): ws[f"{cl}{rr}"].font = F(color=GREEN_LINK if rr in (6, 7) else MUTED, size=9)
        ws[f"{cl}8"] = f"{SECTOR_LABEL[s_]}: {col}"
        ws[f"{cl}8"].font = F(bold=True, size=9); ws[f"{cl}8"].fill = HDR_FILL; ws[f"{cl}8"].border = BOT
        ws[f"{cl}8"].alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center")
        ws.column_dimensions[cl].width = 12
        for i in range(NROWS):
            r = R0 + i
            c = ws[f"{cl}{r}"]
            c.value = f'=IF(${rowcol[tab]}{r}="","",INDEX({tab}!{RAWRNG},${rowcol[tab]}{r},{cl}$6))'
            c.number_format = FMT_M; c.font = F(color=GREEN_LINK)
    ws.column_dimensions["A"].width = 10; ws.column_dimensions["B"].width = 9; ws.column_dimensions["C"].width = 12
    for t in tabs: ws.column_dimensions[rowcol[t]].width = 8
    ws.row_dimensions[8].height = 54
    ws.freeze_panes = ws[f"{L(c0)}{R0}"]
    return ws
PN = build_pull("nsa"); PS = build_pull("sa")

# ---------------------------------------------------------------- Segments
SG = sheet("Segments", "2A78D6")
SG["A1"] = "Segments: building segments as signed sums of Census series (combined private + public, $ millions)"
SG["A1"].font = F(size=13, bold=True)
SG["A2"] = ("Row 7 states each formula in words. Blank cells beyond the latest month are expected. "
            "NSA columns feed YTD, annual and mix figures; SAAR columns feed month-to-month rate figures only.")
SG["A2"].font = F(color=MUTED)
for i, lab in enumerate(["ID", "Name", "Level", "Parent", "Definition", "Adjustment"]):
    c = SG.cell(3 + i, 3, lab); c.font = F(bold=True, color=MUTED); c.alignment = Alignment(horizontal="right")
for j, h in enumerate(["Month", "Year", "Month #"], 1):
    c = SG.cell(8, j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
for i in range(NROWS):
    r = R0 + i
    SG.cell(r, 1, f"=Pull_NSA!A{r}").number_format = "mmm yyyy"
    SG.cell(r, 2, f"=YEAR(A{r})"); SG.cell(r, 3, f"=MONTH(A{r})")
SEGCOL = {}   # (id, adj) -> column letter
ENTS = []     # (id, name, level, parent, components)
for s in CFG["segments"]:
    ENTS.append((s["id"], s["name"], "Segment", "", s["components"]))
for s in CFG["segments"]:
    for ss in s["subsegments"]:
        ENTS.append((ss["id"], ss["name"], "Subsegment", s["name"], ss["components"]))
    if s["subsegments"]:
        resid = list(s["components"]) + [dict(c, sign=-c["sign"]) for ss in s["subsegments"] for c in ss["components"]]
        ENTS.append((s["id"] + "_resid", "Other / not shown separately", "Subsegment", s["name"], resid))
for s in CFG["segments"]:
    for b in s.get("breakdown", []):
        ENTS.append((b["id"], b["name"], "Pie slice", s["name"], b["components"]))

def merged(comps):
    m = {}
    for c in comps:
        k = (src(c), c["col"]); m[k] = m.get(k, 0) + c["sign"]
    return [(k, v) for k, v in m.items() if v]

def defn(comps):
    return " ".join(f'{"+" if v > 0 else "−"} {TAB_OF[(k[0], "nsa")]} "{k[1]}"' for k, v in merged(comps)).lstrip("+ ")

col = 4
def seg_col(eid, nm, level, parent, comps, adj, total_of=None):
    global col
    cl = L(col); col += 1
    SEGCOL[(eid, adj)] = cl
    SG[f"{cl}3"] = eid; SG[f"{cl}4"] = nm; SG[f"{cl}5"] = level; SG[f"{cl}6"] = parent
    SG[f"{cl}7"] = ("Sum of all segments" if total_of else defn(comps)); SG[f"{cl}8"] = nm
    SG[f"{cl}9"] = None
    for rr in range(3, 8):
        SG[f"{cl}{rr}"].font = F(color=MUTED, size=8); SG[f"{cl}{rr}"].alignment = Alignment(wrap_text=True, vertical="top")
    h = SG[f"{cl}8"]; h.font = F(bold=True, size=9); h.fill = HDR_FILL; h.border = BOT
    h.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center")
    if eid in FOCUS: h.fill = fill(shade(FOCUS[eid], .75))
    SG.column_dimensions[cl].width = 13
    pull = "Pull_NSA" if adj == "nsa" else "Pull_SA"
    for i in range(NROWS):
        r = R0 + i
        if total_of:
            expr = "+".join(f"{c}{r}" for c in total_of)
        else:
            parts = []
            for (k, v) in merged(comps):
                ref = f"{pull}!{PULLCOL[(k[0], k[1], adj)]}{r}"
                parts.append(("+" if v > 0 else "-") + (f"{abs(v)}*" if abs(v) != 1 else "") + ref)
            expr = "".join(parts).lstrip("+")
        c = SG[f"{cl}{r}"]
        c.value = f'=IF(A{r}>LatestMonth,"",{expr})'; c.number_format = FMT_M
    SG[f"{cl}9"].number_format = FMT_M
    return cl

# NSA block
seg_ids = [s["id"] for s in CFG["segments"]]
tmp_cols = []
col_total_nsa = col; col += 1   # reserve Total first
for (eid, nm, level, parent, comps) in ENTS:
    seg_col(eid, nm, level, parent, comps, "nsa")
saved = col; col = col_total_nsa
seg_col("total", "Total buildings", "Total", "", [], "nsa", total_of=[SEGCOL[(i, "nsa")] for i in seg_ids])
col = saved + 1
SAhead = L(col - 1)
SG[f"{SAhead}8"] = "SAAR →"; SG[f"{SAhead}8"].font = F(bold=True, color=MUTED); SG.column_dimensions[SAhead].width = 7
col_total_sa = col; col += 1
for s in CFG["segments"]:
    seg_col(s["id"], s["name"], "Segment", "", s["components"], "sa")
for ss_id in ["av_terminal"]:
    e = next(x for x in ENTS if x[0] == ss_id); seg_col(e[0], e[1], e[2], e[3], e[4], "sa")
saved = col; col = col_total_sa
seg_col("total", "Total buildings", "Total", "", [], "sa", total_of=[SEGCOL[(i, "sa")] for i in seg_ids])
col = saved
for cl in [SEGCOL[(e[0], "nsa")] for e in ENTS] + [SEGCOL[("total", "nsa")]]:
    SG[f"{cl}6"].value = SG[f"{cl}6"].value
for (eid, adj), cl in SEGCOL.items():
    SG[f"{cl}8"] = SG[f"{cl}8"].value + ("" if adj == "nsa" else " (SAAR)")
    SG[f"{cl}9"].value = SG[f"{cl}9"].value
SG["C8"].value = "Month #"
SG.row_dimensions[7].height = 60; SG.row_dimensions[8].height = 42
SG.column_dimensions["A"].width = 10; SG.column_dimensions["B"].width = 7; SG.column_dimensions["C"].width = 8
SG.freeze_panes = "D9"
# Re-write rows 3-8 metadata for "Adjustment" row
for (eid, adj), cl in SEGCOL.items():
    pass

YR = f"Segments!$B${R0}:$B${RN}"; MO = f"Segments!$C${R0}:$C${RN}"; MD = f"Segments!$A${R0}:$A${RN}"
def rng(eid, adj="nsa"): cl = SEGCOL[(eid, adj)]; return f"Segments!${cl}${R0}:${cl}${RN}"
def YTD(eid, y): return f"SUMIFS({rng(eid)},{YR},{y},{MO},\"<=\"&LM)"
def ANN(eid, y): return f"SUMIFS({rng(eid)},{YR},{y})"
def MON(eid, d, adj="nsa"): return f"SUMIFS({rng(eid, adj)},{MD},{d})"
NAMES = {e[0]: e[1] for e in ENTS}; NAMES["total"] = "Total buildings"

# ---------------------------------------------------------------- ChartData
CD = sheet("ChartData", "8A8E93")
CD["A1"] = "ChartData: series behind the Dashboard charts (formulas; no typed values)"; CD["A1"].font = F(size=13, bold=True)
# Trend: 120 months ending at the latest month
CD["A3"] = "Trailing 12-month spending, last 120 months ($ millions, NSA)"; CD["A3"].font = F(bold=True)
TR_IDS = ["total", "data_center", "health_care", "aviation", "educational", "av_terminal"]
for j, h in enumerate(["Month"] + [NAMES[i] for i in TR_IDS], 1):
    c = CD.cell(4, j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT; c.alignment = Alignment(wrap_text=True)
TR0 = 5
for i in range(120):
    r = TR0 + i
    CD.cell(r, 1, f"=EDATE(LatestMonth,{i-119})").number_format = "mmm yy"
    for j, eid in enumerate(TR_IDS):
        c = CD.cell(r, 2 + j, f'=SUMIFS({rng(eid)},{MD},">="&EDATE($A{r},-11),{MD},"<="&$A{r})')
        c.number_format = FMT_M
TRN = TR0 + 119
# Mix: total pie
MX = 2 + len(TR_IDS) + 1
mixcol = L(MX)
CD.cell(3, MX, "Total buildings mix, YTD current year ($ millions)").font = F(bold=True)
for j, h in enumerate(["Segment", "YTD", "Highlighted"]):
    c = CD.cell(4, MX + j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
import csv as _csv
# fixed order: focus first, others by latest pipeline YTD size if available
others = [s["id"] for s in CFG["segments"] if s["id"] not in FOCUS]
try:
    dv = {}
    for rrow in _csv.DictReader(open(os.path.join(HERE, "out", "derived.csv"))):
        if rrow["level"] == "segment" and rrow["sector"] == "total" and rrow["adjustment"] == "nsa" and rrow["value"]:
            dv[rrow["segment_id"]] = dv.get(rrow["segment_id"], 0) + float(rrow["value"])
    others.sort(key=lambda i: -dv.get(i, 0))
except Exception:
    pass
MIX_IDS = ["data_center", "health_care", "aviation", "educational"] + others
for i, eid in enumerate(MIX_IDS):
    r = 5 + i
    CD.cell(r, MX, NAMES[eid])
    CD.cell(r, MX + 1, "=" + YTD(eid, "LY")).number_format = FMT_M
    CD.cell(r, MX + 2, "Yes" if eid in FOCUS else "")
MIX_N = len(MIX_IDS)
# Breakdown pies
BD = {}
bc = MX + 4
for s in CFG["segments"]:
    if s["id"] not in FOCUS or not s.get("breakdown"): continue
    CD.cell(3, bc, f"{s['name']} by subcategory, YTD ($ millions)").font = F(bold=True)
    for j, h in enumerate(["Subcategory", "Sector", "YTD", "Plotted (≥0)"]):
        c = CD.cell(4, bc + j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
    rows = []
    for i, b in enumerate(s["breakdown"]):
        r = 5 + i
        pub = all(c_["sector"] == "public" for c_ in b["components"])
        label = b["name"] if not pub else ("Public: " + b["name"].replace(" (state & local)", " (S&L)"))
        CD.cell(r, bc, label)
        CD.cell(r, bc + 1, "Public" if pub else "Private")
        CD.cell(r, bc + 2, "=" + YTD(b["id"], "LY")).number_format = FMT_M
        CD.cell(r, bc + 3, f"=MAX(0,{L(bc+2)}{r})").number_format = FMT_M
        rows.append((r, pub))
    BD[s["id"]] = (bc, rows)
    for j, w in enumerate([30, 9, 11, 11]): CD.column_dimensions[L(bc + j)].width = w
    bc += 5
CD.column_dimensions["A"].width = 9
for j in range(len(TR_IDS)): CD.column_dimensions[L(2 + j)].width = 13
CD.column_dimensions[L(MX)].width = 28; CD.column_dimensions[L(MX + 1)].width = 11

# ---------------------------------------------------------------- Dashboard
DB = wb.create_sheet("Dashboard", 0); DB.sheet_properties.tabColor = NAVY; DB.sheet_view.showGridLines = False
DB.column_dimensions["A"].width = 1.5
DB.column_dimensions["B"].width = 38
for c_ in "CDEFGHIJ": DB.column_dimensions[c_].width = 13
for c_ in "KLMNOP": DB.column_dimensions[c_].width = 11
DB["B1"] = "Nonresidential Building Construction Spending"; DB["B1"].font = F(size=18, bold=True, color=INK)
DB["B2"] = ("US Census Bureau, Value of Construction Put in Place (C30). Buildings only: civil, water, power, communication and "
            "transportation infrastructure removed. Private + public combined. Nominal dollars.")
DB["B2"].font = F(color=MUTED)
DB["B3"] = '="Census release: "&ReleaseText&"   ·   Data through "&TEXT(LatestMonth,"mmm yyyy")&" ("&LatestFlag&")   ·   "&Checks!B' + str(rS)
DB["B3"].font = F(bold=True, color=INK)
DB.conditional_formatting.add("B3", FormulaRule(formula=[f'COUNTIF(Checks!$G$5:$H${rS-2},"MISMATCH")>0'], font=F(color="B42F2F", bold=True)))

def section(r, title, note=None):
    DB[f"B{r}"] = title; DB[f"B{r}"].font = F(size=13, bold=True, color=NAVY)
    for c_ in "BCDEFGHIJ": DB[f"{c_}{r}"].border = Border(bottom=Side(style="medium", color=NAVY))
    if note:
        DB[f"B{r+1}"] = note; DB[f"B{r+1}"].font = F(color=MUTED, size=9)

KPI_HEAD = ['="YTD "&LY', '="YTD "&(LY-1)', '="Change"', '="YTD growth"',
            '=TEXT(LatestMonth,"mmm")&" vs "&TEXT(LatestMonth,"mmm")&" "&(LY-1)',
            '="SAAR "&TEXT(LatestMonth,"mmm")', '="SAAR vs prior month"', '="Share of buildings"']
def table_head(r, first="Segment"):
    DB[f"B{r}"] = first
    for j, h in enumerate(KPI_HEAD):
        DB.cell(r, 3 + j, h)
    for c_ in "BCDEFGHIJ":
        c = DB[f"{c_}{r}"]; c.font = F(bold=True, color=INK, size=9); c.fill = HDR_FILL; c.border = BOT
        c.alignment = Alignment(horizontal="left" if c_ == "B" else "right", wrap_text=True, vertical="bottom")
    DB.row_dimensions[r].height = 28

TOTAL_YTD = {}
def kpi_row(r, eid, label, style="seg", has_sa=None):
    DB[f"B{r}"] = label
    DB[f"C{r}"] = "=" + YTD(eid, "LY")
    DB[f"D{r}"] = "=" + YTD(eid, "LY-1")
    DB[f"E{r}"] = f"=C{r}-D{r}"
    DB[f"F{r}"] = f'=IF(D{r}>0,C{r}/D{r}-1,"")'
    DB[f"G{r}"] = f'=IF({MON(eid,"EDATE(LatestMonth,-12)")}>0,{MON(eid,"LatestMonth")}/{MON(eid,"EDATE(LatestMonth,-12)")}-1,"")'
    if (eid, "sa") in SEGCOL:
        DB[f"H{r}"] = "=" + MON(eid, "LatestMonth", "sa")
        DB[f"I{r}"] = f'=IF({MON(eid,"EDATE(LatestMonth,-1)","sa")}>0,H{r}/{MON(eid,"EDATE(LatestMonth,-1)","sa")}-1,"")'
    else:
        DB[f"H{r}"] = "—"; DB[f"I{r}"] = "—"
    DB[f"J{r}"] = f"=IF($C${TOTAL_YTD['row']}>0,C{r}/$C${TOTAL_YTD['row']},\"\")" if eid != "total" else "=1"
    for c_, fmt in zip("CDEFGHIJ", [FMT_B, FMT_B, FMT_DB, FMT_P, FMT_P, FMT_B, FMT_P, FMT_SH]):
        DB[f"{c_}{r}"].number_format = fmt
        DB[f"{c_}{r}"].alignment = Alignment(horizontal="right")
    for c_ in "BCDEFGHIJ":
        c = DB[f"{c_}{r}"]; c.border = BOT
        if style == "total": c.font = F(bold=True); c.fill = fill("F4F5F2")
        elif style == "sub": c.font = F(color=MUTED, size=9, italic=(eid.endswith("_resid")))
        else: c.font = F(bold=(eid in FOCUS))
    if style == "sub": DB[f"B{r}"].alignment = Alignment(indent=2)
    if eid in FOCUS:
        DB[f"A{r}"].fill = fill(FOCUS[eid])
    if eid == "av_terminal":
        DB[f"A{r}"].fill = fill(shade(FOCUS["aviation"], .45))

# --- Current year
r = 5
section(r, "Current year", "Year-to-date = sum of not-seasonally-adjusted months, January through the latest month. "
        "SAAR = seasonally adjusted annual rate; month-to-month only. $B = billions.")
table_head(r + 2)
TOTAL_YTD["row"] = r + 3
kpi_row(r + 3, "total", "Total nonresidential buildings", "total")
kpi_rows = [("data_center", "Data centers"), ("health_care", "Health care"), ("aviation", "Aviation"),
            ("av_terminal", "of which: Air passenger terminals (state & local)"), ("educational", "Educational")]
for i, (eid, lab) in enumerate(kpi_rows):
    kpi_row(r + 4 + i, eid, lab, "sub" if eid == "av_terminal" else "seg")
CY_END = r + 4 + len(kpi_rows) - 1
gr = f"F{r+3}:G{CY_END} I{r+3}:I{CY_END}"

# --- Spending mix (pies)
r = CY_END + 2
section(r, "Spending mix, year to date",
        "Left: total buildings by segment (highlighted segments in colour, others in grey). Right: each highlighted segment by "
        "Census subcategory; patterned slices are public spending. Public detail comes from the state-and-local tabs; Federal = public total − state & local.")
MIX_ROW = r + 3

def pie(title, cats, vals, colors, patterns=None, w=12.5, h=8.6, legend="r"):
    pc = PieChart()
    pc.title = title
    pc.add_data(vals, titles_from_data=False)
    pc.set_categories(cats)
    pc.width, pc.height = w, h
    s = pc.series[0]
    s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="FFFFFF", w=12700))
    for i, colr in enumerate(colors):
        pt = DataPoint(idx=i)
        if patterns and patterns[i]:
            pt.graphicalProperties = GraphicalProperties(
                pattFill=PatternFillProperties(prst="wdUpDiag", fgClr=ColorChoice(srgbClr=colr), bgClr=ColorChoice(srgbClr="FFFFFF")),
                ln=LineProperties(solidFill="FFFFFF", w=12700))
        else:
            pt.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill="FFFFFF", w=12700))
        s.dPt.append(pt)
    s.dLbls = DataLabelList(); s.dLbls.showPercent = True; s.dLbls.showVal = False; s.dLbls.showCatName = False
    s.dLbls.showSerName = False; s.dLbls.showLeaderLines = True; s.dLbls.position = "bestFit"
    pc.legend.position = legend
    return pc

cats = Reference(CD, min_col=MX, min_row=5, max_row=4 + MIX_N)
vals = Reference(CD, min_col=MX + 1, min_row=5, max_row=4 + MIX_N)
gi = iter(GRAYS)
colors = [FOCUS[e] if e in FOCUS else next(gi) for e in MIX_IDS]
DB.add_chart(pie("Total buildings by segment", cats, vals, colors, w=15.2, h=19.6), f"B{MIX_ROW}")
anchors = [f"F{MIX_ROW}", f"J{MIX_ROW}"]
pos = 0
pie_anchor_rows = [MIX_ROW, MIX_ROW + 20]
pie_slots = [("F", MIX_ROW), ("K", MIX_ROW), ("F", MIX_ROW + 19), ("K", MIX_ROW + 19)]
# total pie occupies B..E rows MIX_ROW..+19; four segment pies in a 2x2 grid to its right
for k, sid in enumerate(["data_center", "health_care", "aviation", "educational"]):
    if sid not in BD: continue
    bcol, rows = BD[sid]
    base = FOCUS[sid]
    priv = [i for i, (rr, pub) in enumerate(rows) if not pub]
    cols_, pats = [], []
    pi = pu = 0
    for i, (rr, pub) in enumerate(rows):
        if pub: cols_.append(shade(base, min(.15 * pu, .6))); pats.append(True); pu += 1
        else: cols_.append(shade(base, min(.22 * pi, .66))); pats.append(False); pi += 1
    cats = Reference(CD, min_col=bcol, min_row=rows[0][0], max_row=rows[-1][0])
    vals = Reference(CD, min_col=bcol + 3, min_row=rows[0][0], max_row=rows[-1][0])
    nm = next(s["name"] for s in CFG["segments"] if s["id"] == sid)
    cl_, rw_ = pie_slots[k]
    DB.add_chart(pie(nm, cats, vals, cols_, pats, w=12.6, h=9.6, legend="b"), f"{cl_}{rw_}")
MIX_END = MIX_ROW + 38

# --- Ten-year trend charts
r = MIX_END + 1
section(r, "Ten-year trend", "Trailing 12-month spending (sum of the last 12 not-seasonally-adjusted months), last 120 months. "
        "The window rolls forward automatically when new data is pasted.")
TREND_ROW = r + 3
def line(title, ids, colors, w, h, dash_first=False):
    lc = LineChart(); lc.title = title; lc.width, lc.height = w, h
    lc.y_axis.title = "$ millions"; lc.y_axis.numFmt = '#,##0'; lc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E2E4DF"))
    lc.x_axis = DateAxis(crossAx=100); lc.x_axis.number_format = "yyyy"
    lc.x_axis.majorTimeUnit = "years"; lc.x_axis.majorUnit = 1; lc.x_axis.baseTimeUnit = "months"
    lc.x_axis.delete = False; lc.y_axis.delete = False; lc.y_axis.crossAx = 500
    for eid, colr in zip(ids, colors):
        j = TR_IDS.index(eid)
        data = Reference(CD, min_col=2 + j, min_row=4, max_row=TRN)
        lc.add_data(data, titles_from_data=True)
        s = lc.series[-1]; s.smooth = False
        s.graphicalProperties.line.solidFill = colr; s.graphicalProperties.line.width = 28575
        s.marker.symbol = "none"
    lc.set_categories(Reference(CD, min_col=1, min_row=TR0, max_row=TRN))
    lc.legend.position = "b"
    return lc
DB.add_chart(line("Highlighted segments: trailing 12-month spending",
                  ["data_center", "health_care", "aviation", "educational", "av_terminal"],
                  [FOCUS["data_center"], FOCUS["health_care"], FOCUS["aviation"], FOCUS["educational"], shade(FOCUS["aviation"], .45)],
                  w=24, h=10), f"B{TREND_ROW}")
DB.add_chart(line("Total nonresidential buildings: trailing 12-month spending", ["total"], ["35383C"], w=16, h=10), f"H{TREND_ROW}")
TREND_END = TREND_ROW + 21

# --- YTD by segment (all segments + subsegments)
r = TREND_END + 1
section(r, "Year to date by segment",
        "Subcategories are indented. \"Other / not shown separately\" = segment minus its listed subcategories (Census does not publish every subcategory).")
table_head(r + 2)
kpi_row(r + 3, "total", "Total nonresidential buildings", "total")
rr = r + 4
order = ["data_center", "health_care", "aviation", "educational"] + others
for sid in order:
    s = next(x for x in CFG["segments"] if x["id"] == sid)
    kpi_row(rr, sid, s["name"]); rr += 1
    for ss in s["subsegments"]:
        kpi_row(rr, ss["id"], ss["name"], "sub"); rr += 1
    if s["subsegments"]:
        kpi_row(rr, sid + "_resid", "Other / not shown separately", "sub"); rr += 1
SEGT_END = rr - 1
gr += f" F{r+3}:G{SEGT_END} I{r+3}:I{SEGT_END}"

# --- Annual history
r = SEGT_END + 2
section(r, "Annual history ($ billions)", "Calendar-year sums of not-seasonally-adjusted months. CAGR runs from the first to the last full year shown.")
hr = r + 2
DB[f"B{hr}"] = "Segment"
ycols = [L(3 + i) for i in range(10)]
for i, c_ in enumerate(ycols):
    DB[f"{c_}{hr}"] = f"=FirstFullYear+{i}"
ytdc, ytdpc, cagc = L(13), L(14), L(15)
DB[f"{ytdc}{hr}"] = '="YTD "&LY'; DB[f"{ytdpc}{hr}"] = '="YTD "&(LY-1)'; DB[f"{cagc}{hr}"] = "CAGR"
for c_ in ["B"] + ycols + [ytdc, ytdpc, cagc]:
    c = DB[f"{c_}{hr}"]; c.font = F(bold=True, size=9); c.fill = HDR_FILL; c.border = BOT
    c.alignment = Alignment(horizontal="left" if c_ == "B" else "right"); c.number_format = "0"
for c_ in "KLMNOP": DB[f"{c_}{r}"].border = Border(bottom=Side(style="medium", color=NAVY))
ar = hr + 1
for eid in ["total"] + order:
    lab = "Total nonresidential buildings" if eid == "total" else NAMES[eid]
    DB[f"B{ar}"] = lab
    for c_ in ycols:
        DB[f"{c_}{ar}"] = "=" + ANN(eid, f"{c_}${hr}")
    DB[f"{ytdc}{ar}"] = "=" + YTD(eid, "LY"); DB[f"{ytdpc}{ar}"] = "=" + YTD(eid, "LY-1")
    DB[f"{cagc}{ar}"] = f'=IF({ycols[0]}{ar}>0,({ycols[-1]}{ar}/{ycols[0]}{ar})^(1/9)-1,"")'
    for c_ in ycols + [ytdc, ytdpc]: DB[f"{c_}{ar}"].number_format = '#,##0.0,;(#,##0.0,);"-"'
    DB[f"{cagc}{ar}"].number_format = FMT_P
    for c_ in ["B"] + ycols + [ytdc, ytdpc, cagc]:
        c = DB[f"{c_}{ar}"]; c.border = BOT
        c.font = F(bold=(eid == "total" or eid in FOCUS))
        if eid == "total": c.fill = fill("F4F5F2")
    if eid in FOCUS: DB[f"A{ar}"].fill = fill(FOCUS[eid])
    ar += 1
gr += f" {cagc}{hr+1}:{cagc}{ar-1}"
ANN_END = ar - 1
DB[f"B{ANN_END+2}"] = ("Every figure on this sheet is a formula. Trace any cell with Formulas → Trace Precedents: Dashboard → Segments → "
                       "Pull_NSA / Pull_SA → raw Census tab (row from the month, column from Map).")
DB[f"B{ANN_END+2}"].font = F(color=MUTED, size=9, italic=True)

for g in gr.split():
    DB.conditional_formatting.add(g, CellIsRule(operator="greaterThan", formula=["0.0005"], font=F(color="006300")))
    DB.conditional_formatting.add(g, CellIsRule(operator="lessThan", formula=["-0.0005"], font=F(color="B42F2F")))
DB.freeze_panes = "C4"
DB.sheet_view.zoomScale = 90

# ---------------------------------------------------------------- Checks: reconciliation block
r = CK_RECON
CK.cell(r, 1, "Reconciliation, latest month (not seasonally adjusted, $ millions)").font = F(size=12, bold=True)
for j, h in enumerate(["Check", "What it confirms", "Reference", "Computed", "Difference", "Status"], 1):
    c = CK.cell(r + 1, j, h); c.font = F(bold=True); c.fill = HDR_FILL; c.border = BOT
rr = r + 2
def ck(label, what, ref, comp, tol):
    global rr
    CK.cell(rr, 1, label); CK.cell(rr, 2, what).font = F(color=MUTED, size=9)
    CK.cell(rr, 3, ref).number_format = FMT_M; CK.cell(rr, 4, comp).number_format = FMT_M
    CK.cell(rr, 5, f"=C{rr}-D{rr}").number_format = FMT_M
    CK.cell(rr, 6, f'=IF(ABS(E{rr})<={tol},"OK","CHECK")' if tol is not None else "Info")
    rr += 1
for s in CFG["segments"]:
    if s.get("breakdown"):
        ck(f"{s['name']}: pie slices", "Slices add up to the segment",
           "=" + MON(s["id"], "LatestMonth"), "=" + "+".join(MON(b["id"], "LatestMonth") for b in s["breakdown"]), 0)
hc = next(s for s in CFG["segments"] if s["id"] == "health_care")
ck("Health care: subcategories", "Listed subcategories vs parent (Census rounding)",
   "=" + MON("health_care", "LatestMonth"), "=" + "+".join(MON(x["id"], "LatestMonth") for x in hc["subsegments"]), 2)
def pullv(s_, col, adj="nsa"):
    return f"SUMIFS(Pull_{adj.upper()}!${PULLCOL[(s_, col, adj)]}${R0}:${PULLCOL[(s_, col, adj)]}${RN},Pull_NSA!$A${R0}:$A${RN},LatestMonth)"
for cat, slc in [("Health care", "Health care"), ("Educational", "Educational")]:
    ck(f"Federal {cat.lower()}", "Public − state & local = Census Federal column",
       "=" + pullv("public", f"Federal {cat}"), f"={pullv('public', 'Public ' + cat)}-{pullv('statelocal', slc)}", 2)
priv_comps = [k for k in SEG_COMPS if k[0] == "private"]
ck("Private buildings share", "Census private nonresidential; difference = excluded infrastructure",
   "=" + pullv("private", "Nonresidential"),
   "=" + "+".join(f"{'' if v>0 else '-'}{pullv(k[0], k[1])}" for s in CFG["segments"] for (k, v) in merged(s["components"]) if k[0] == "private").replace("+-", "-"), None)
ck("Public buildings share", "Census public nonresidential; difference = excluded infrastructure",
   "=" + pullv("public", "Public Nonresidential"),
   "=" + "+".join(f"{'' if v>0 else '-'}{pullv(k[0], k[1])}" for s in CFG["segments"] for (k, v) in merged(s["components"]) if k[0] in ("public", "statelocal")).replace("+-", "-"), None)
CK.conditional_formatting.add(f"F{r+2}:F{rr}", CellIsRule(operator="equal", formula=['"CHECK"'], font=F(color="B42F2F", bold=True)))
CK.conditional_formatting.add(f"F{r+2}:F{rr}", CellIsRule(operator="equal", formula=['"OK"'], font=F(color="006300", bold=True)))

# ---------------------------------------------------------------- README
RM = wb.create_sheet("README", 1); RM.sheet_properties.tabColor = "1F2328"; RM.sheet_view.showGridLines = False
RM.column_dimensions["A"].width = 2; RM.column_dimensions["B"].width = 26; RM.column_dimensions["C"].width = 110
lines = [
    ("h", "C30 Building Dashboard: how this workbook works"),
    ("p", "Source: U.S. Census Bureau, Value of Construction Put in Place (C30), historical time-series workbooks. "
          "https://www.census.gov/construction/c30/historical_data.html"),
    ("p", "All values are nominal dollars, millions as published (the Dashboard displays billions). Year-to-date and annual figures sum "
          "not-seasonally-adjusted months, as Census does; SAAR values are used only for month-level rates and are never summed."),
    ("h2", "Monthly update (copy/paste)"),
    ("n", "1. Download the six workbooks from the Census page above: privtime, pubtime, sltime, privsatime, pubsatime, slsatime (.xlsx)."),
    ("n", "2. In each downloaded file, click the top-left corner of the sheet to select all, and copy."),
    ("n", "3. In this workbook, click cell A1 of the tab with the same name and paste (Paste Values is fine). Paste over the whole tab; do not insert rows."),
    ("n", "4. Open the Checks tab. The Status cell must read \"All raw tabs consistent\". A MISMATCH means one tab is from a different release or month."),
    ("n", "5. The Dashboard, trend window, annual table and pies update automatically. Recalculate (F9) if calculation is set to manual."),
    ("p", "Linking instead of pasting: each raw tab can instead hold external links to the downloaded files "
          "(e.g. ='[privtime.xlsx]Private NSA'!A1 filled across A1:BZ500). Lookups find months and columns by label, so either works."),
    ("h2", "Tabs (calculation flow, top to bottom)"),
    ("t", "Dashboard", "Current-year KPIs, spending mix pies, 10-year trend charts, YTD table for every segment and subcategory, annual history."),
    ("t", "Checks", "Latest month and release read from each raw tab, consistency status, named control cells, and reconciliation checks."),
    ("t", "Segments", "Each building segment, subcategory and pie slice as a signed sum of Census series (definition stated in row 7). Combined private + public."),
    ("t", "Pull_NSA / Pull_SA", "One column per Census series, one row per month. Row numbers per raw tab and the column letter from Map give the exact source cell."),
    ("t", "Map", "Locates each Census series in its raw tab by header label. Edit the blue pattern if Census renames a header."),
    ("t", "ChartData", "Formula-driven series behind the charts (trailing 12-month trend, mix tables)."),
    ("t", "privtime … slsatime", "Raw Census sheets, pasted unchanged. These are the only inputs."),
    ("h2", "Segment definitions"),
]
for s in CFG["segments"]:
    lines.append(("t", s["name"] + (" (highlighted)" if s["id"] in FOCUS else ""), defn(s["components"]) + (f". {s['note']}" if s.get("note") else "")))
lines.append(("h2", "Excluded from buildings"))
for e in CFG["excluded"]:
    lines.append(("t", e["col"], f"{e['sector'].capitalize()}: {e['reason']}"))
lines += [("h2", "Colour key"),
          ("p", "Green text = value pulled from a raw tab. Blue text = editable lookup pattern (Map). Black = calculation. "
                "Coloured stripe in Dashboard column A = highlighted segment (data centers blue, health care orange, aviation green, education violet).")]
r = 2
for ln in lines:
    kind = ln[0]
    if kind == "h": RM[f"B{r}"] = ln[1]; RM[f"B{r}"].font = F(size=16, bold=True); r += 1
    elif kind == "h2":
        r += 1; RM[f"B{r}"] = ln[1]; RM[f"B{r}"].font = F(size=12, bold=True, color=NAVY)
        for c_ in "BC": RM[f"{c_}{r}"].border = Border(bottom=Side(style="thin", color=NAVY))
    elif kind in ("p", "n"):
        RM[f"B{r}"] = ln[1]; RM.merge_cells(f"B{r}:C{r}")
        RM[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top"); RM.row_dimensions[r].height = 30 if len(ln[1]) > 120 else 16
    elif kind == "t":
        RM[f"B{r}"] = ln[1]; RM[f"B{r}"].font = F(bold=True); RM[f"B{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        RM[f"C{r}"] = ln[2]; RM[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        RM.row_dimensions[r].height = 16 * max(1, (len(ln[2]) // 105) + 1)
    r += 1

wb.remove(first)
order = ["Dashboard", "README", "Checks", "Segments", "Pull_NSA", "Pull_SA", "Map", "ChartData"] + [t for t, _, _ in TABS]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.calculation.fullCalcOnLoad = True
os.makedirs(os.path.dirname(OUTF), exist_ok=True)
wb.save(OUTF)
print("wrote", OUTF, "| components", len(comp_keys), "| segment columns", len(SEGCOL))
